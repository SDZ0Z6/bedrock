"""生命周期标签（正常、结算、风控……）。

可选的标签放在台账第二个工作表 LIFECYCLE（NAME、COLOR 两列）里；没有这张表就是默认的
正常 / 结算 / 风控，第一次在页面上增删标签时才把表建出来。账号身上的标签在 LIFECYCLE 列，
一格多个（读的时候中英文逗号、顿号、分号、竖线、换行都认，写回统一成半角逗号）。

页面上三处能动它：表格里每行的铅笔（整页共用一个小弹层，fetch 调 /accounts/lifecycle，
回 JSON，不整页刷新）、右上角「生命周期标签」弹窗（增删清单）、新增 / 修改弹窗里的多选框。
只是标记，不影响查询和告警，所以这些操作都不备份台账，只记审计日志。
"""

from __future__ import annotations

import re

import openpyxl
import pytest
from werkzeug.datastructures import MultiDict

from bedrock_cost import config, excel_source
from bedrock_cost.excel_source import ExcelSourceError, LedgerConflict, LifecycleTag, load_accounts

from .conftest import LEDGER_ROWS, write_ledger
from .test_accounts import (  # noqa: F401
    EMAILS,
    NEW_FORM,
    _dialog,
    _edit,
    admin,
    by_account,
    input_attrs,
    inputs,
    page,
    post,
    raw_rows,
    reopened_dialogs,
    rewrite_ledger,
    table_rows,
    text,
    toasts,
    token,
)

FETCH = {"X-Requested-With": "fetch"}
COLOR = {key: hex_ for key, (_, hex_) in excel_source.LIFECYCLE_COLORS.items()}
GREEN, GRAY, RED, AMBER = COLOR["green"], COLOR["gray"], COLOR["red"], COLOR["amber"]
DEFAULTS = [LifecycleTag("正常", "green"), LifecycleTag("结算", "gray"), LifecycleTag("风控", "red")]
KEY_ONE = "111111111111#2"          # conftest 台账里第一个账号的 key（账号#行号）
THIRTEEN = "一二三四五六七八九十一二三"   # 比 MAX_LIFECYCLE_NAME 多一个字


# ------------------------------------------------------------------ 帮手
def write_tag_sheet(path, rows) -> None:
    """给台账加上第二个工作表 LIFECYCLE（表头 + rows）。"""
    book = openpyxl.load_workbook(path)
    sheet = book.create_sheet(excel_source.LIFECYCLE_SHEET)
    sheet.append(["NAME", "COLOR"])
    for row in rows:
        sheet.append(list(row))
    book.save(path)
    excel_source.clear_cache()


def tag_sheet(path):
    """直接看文件里的 LIFECYCLE 表（不含表头），[[名字, 颜色], …]；还没有这张表就是 None。"""
    book = openpyxl.load_workbook(path, data_only=True)
    try:
        if excel_source.LIFECYCLE_SHEET not in book.sheetnames:
            return None
        return [list(row) for row in book[excel_source.LIFECYCLE_SHEET].iter_rows(min_row=2, values_only=True)]
    finally:
        book.close()


def life_cell(path, row: int = 0):
    """账号那一行 LIFECYCLE 格里存的原值（没有这一列就是 None）。"""
    header, *rows = raw_rows(path)
    return rows[row][header.index("LIFECYCLE")] if "LIFECYCLE" in header else None


def backups(ledger) -> list:
    return sorted((ledger.parent / excel_source.BACKUP_DIR_NAME).glob("*.xlsx"))


def audit(ledger) -> str:
    path = ledger.parent / excel_source.AUDIT_NAME
    return path.read_text(encoding="utf-8") if path.exists() else ""


def emails(**more) -> dict:
    """rewrite_ledger 用的列：两个账号的邮箱（admin 台账本来就有），再加几列。"""
    return {"EMAIL": [EMAILS["111111111111"], EMAILS["222222222222"]], **more}


def set_life(client, key: str, *tags: str, fetch: bool = True):
    """表格里的小弹层点「保存」：fetch=True 是页面脚本发的，False 是没开 JS 的普通提交。"""
    return client.post(
        "/accounts/lifecycle", data={"csrf": token(client), "key": key, "lifecycle": list(tags)},
        headers=FETCH if fetch else {},
    )


def offered_tags(fragment: str) -> list[str]:
    return [field["value"] for field in inputs(fragment, "lifecycle")]


def checked_tags(fragment: str) -> list[str]:
    return [field["value"] for field in inputs(fragment, "lifecycle") if field.get("checked")]


def tag_dialog(html: str) -> str:
    return _dialog(html, 'id="dlg-life"')


def add_form(html: str) -> str:
    dialog = tag_dialog(html)
    return dialog[dialog.index('action="/accounts/lifecycle/add"') :]


def listed_tags(html: str) -> dict[str, int]:
    """「生命周期标签」弹窗里列出的标签和各自的账号数，按清单顺序。"""
    rows = re.findall(r'<li class="life-row">(.*?)</li>', tag_dialog(html), re.S)
    return {name: int(count) for row in rows for name, count in re.findall(r'data-life-count="([^"]+)">(\d+)<', row)}


def life_error(html: str) -> str:
    found = re.search(r'class="[^"]*\blife-error\b[^"]*"[^>]*>(.*?)</div>', tag_dialog(html), re.S)
    return text(found.group(1)) if found else ""


def filter_chips(html: str) -> list[str]:
    chips = html[html.index('class="table-tools tools-life"') : html.index('<div class="table-scroll">')]
    return re.findall(r'<button class="fchip" type="button" data-value="([^"]*)"', chips)


# ------------------------------------------------------------------ 拆分
class TestSplitting:
    @pytest.mark.parametrize(
        "raw, expected",
        [
            (None, ()),
            ("", ()),
            ("正常", ("正常",)),
            ("正常,结算", ("正常", "结算")),
            ("正常，结算", ("正常", "结算")),                 # 全角逗号
            ("正常、结算", ("正常", "结算")),                 # 顿号
            ("正常;结算；风控", ("正常", "结算", "风控")),      # 半角、全角分号
            ("正常|结算", ("正常", "结算")),
            ("正常\n结算", ("正常", "结算")),                 # Excel 里 Alt+Enter 换行
            (" 正常 , ,结算,正常 ", ("正常", "结算")),          # 去空、去重、保持顺序
            ("观察 中", ("观察 中",)),                        # 空格不是分隔符：标签名里可以有空格
        ],
    )
    def test_split(self, raw, expected):
        assert excel_source.split_lifecycle(raw) == expected

    def test_form_ticks_are_merged(self):
        form = MultiDict([("lifecycle", "结算"), ("lifecycle", "正常，风控"), ("lifecycle", "结算"), ("lifecycle", "")])
        assert excel_source.form_lifecycle(form) == ["结算", "正常", "风控"]

    def test_a_plain_dict_works_too(self):
        assert excel_source.form_lifecycle({"lifecycle": "正常、结算"}) == ["正常", "结算"]
        assert excel_source.form_lifecycle({}) == []


# ------------------------------------------------------------------ 读
class TestReading:
    def test_defaults_without_the_second_sheet(self, ledger):
        assert excel_source.load_lifecycle(force=True) == DEFAULTS
        assert excel_source.lifecycle_colors() == {"正常": GREEN, "结算": GRAY, "风控": RED}

    def test_the_second_sheet_is_the_list(self, ledger):
        write_tag_sheet(ledger, [
            ("观察", "Amber"),
            ("正常", " green "),
            (None, "red"),            # 没名字的行跳过
            ("观察", "blue"),          # 重名的只认第一个
            ("老客户", "pink"),        # 认不出的颜色
            ("  续约  ", None),        # 没填颜色
        ])
        assert excel_source.load_lifecycle(force=True) == [
            LifecycleTag("观察", "amber"),
            LifecycleTag("正常", "green"),
            LifecycleTag("老客户", "gray"),
            LifecycleTag("续约", "gray"),
        ]

    def test_an_empty_sheet_is_an_empty_list(self, ledger):
        """表在、一个标签都没有：那就是删光了，不是回到默认的三个。"""
        write_tag_sheet(ledger, [])
        assert excel_source.load_lifecycle(force=True) == []

    def test_edits_to_the_sheet_show_up_without_a_restart(self, ledger):
        """和账号共用一次文件读取、同一份按 mtime+size 的缓存。"""
        assert excel_source.load_lifecycle() == DEFAULTS
        book = openpyxl.load_workbook(ledger)
        sheet = book.create_sheet(excel_source.LIFECYCLE_SHEET)
        sheet.append(["NAME", "COLOR"])
        sheet.append(["观察", "amber"])
        book.save(ledger)
        assert excel_source.load_lifecycle() == [LifecycleTag("观察", "amber")]

    def test_hex(self):
        assert LifecycleTag("甲", "teal").hex == COLOR["teal"]
        assert LifecycleTag("甲", "nope").hex == GRAY

    def test_colors_for_a_given_list(self):
        assert excel_source.lifecycle_colors([LifecycleTag("甲", "blue")]) == {"甲": COLOR["blue"]}
        assert excel_source.lifecycle_colors([]) == {}

    def test_the_accounts_column(self, ledger):
        rewrite_ledger(ledger, LIFECYCLE=["正常，风控、正常", None])
        first, second = load_accounts(force=True)
        assert (first.lifecycle, second.lifecycle) == (("正常", "风控"), ())

    def test_an_old_ledger_has_no_tags_anywhere(self, accounts):
        assert all(account.lifecycle == () for account in accounts)


# ------------------------------------------------------------------ 写：一个账号的标签
class TestSetLifecycle:
    """set_lifecycle：只写这个账号的 LIFECYCLE 一格；不认清单（清单由页面那一层管）。"""

    def test_writes_one_cell(self, ledger):
        note = excel_source.set_lifecycle(KEY_ONE, ["正常", "风控"], actor="tester")
        assert note == "账号 111111111111 的生命周期：空 → 正常、风控"
        assert (life_cell(ledger), life_cell(ledger, 1)) == ("正常,风控", None)
        assert by_account("111111111111").lifecycle == ("正常", "风控")
        assert f"tester\t{note}" in audit(ledger)
        assert backups(ledger) == []                                  # 只是标记，不备份

    def test_the_names_are_cleaned(self, ledger):
        excel_source.set_lifecycle(KEY_ONE, [" 正常 ", "", "正常", "结算"])
        assert life_cell(ledger) == "正常,结算"

    def test_the_same_tags_again_write_nothing(self, ledger):
        excel_source.set_lifecycle(KEY_ONE, ["正常"])
        before = ledger.stat().st_mtime_ns
        assert excel_source.set_lifecycle(KEY_ONE, ["正常"]) == ""
        assert ledger.stat().st_mtime_ns == before

    def test_clearing(self, ledger):
        excel_source.set_lifecycle(KEY_ONE, ["正常"])
        assert excel_source.set_lifecycle(KEY_ONE, []) == "账号 111111111111 的生命周期：正常 → 空"
        assert life_cell(ledger) is None

    def test_names_are_at_most_twelve_characters(self, ledger):
        before = ledger.stat().st_mtime_ns
        with pytest.raises(ExcelSourceError, match="最多 12 个字"):
            excel_source.set_lifecycle(KEY_ONE, ["正常", THIRTEEN])
        assert ledger.stat().st_mtime_ns == before
        assert excel_source.set_lifecycle(KEY_ONE, [THIRTEEN[:12]])

    def test_an_old_ledger_gets_the_column(self, ledger):
        assert "LIFECYCLE" not in raw_rows(ledger)[0]
        excel_source.set_lifecycle(KEY_ONE, ["正常"])
        assert "LIFECYCLE" in raw_rows(ledger)[0]

    def test_a_reshuffled_ledger_is_a_conflict(self, ledger):
        write_ledger(ledger, rows=list(reversed(LEDGER_ROWS)))
        excel_source.clear_cache()
        with pytest.raises(LedgerConflict):
            excel_source.set_lifecycle(KEY_ONE, ["正常"])
        assert life_cell(ledger) is None


# ------------------------------------------------------------------ 写：标签清单
class TestAddLifecycle:
    def test_the_first_add_creates_the_sheet_with_the_defaults(self, ledger):
        assert tag_sheet(ledger) is None
        note = excel_source.add_lifecycle("观察", "amber", actor="tester")
        assert note == "新增生命周期标签「观察」（琥珀色）"
        assert tag_sheet(ledger) == [["正常", "green"], ["结算", "gray"], ["风控", "red"], ["观察", "amber"]]
        assert excel_source.load_lifecycle(force=True) == [*DEFAULTS, LifecycleTag("观察", "amber")]
        assert f"tester\t{note}" in audit(ledger)
        assert backups(ledger) == []

    def test_the_accounts_stay_on_the_first_sheet(self, ledger):
        excel_source.add_lifecycle("观察", "amber")
        assert raw_rows(ledger)[0][0] == "PARTNER"
        assert [a.account for a in load_accounts(force=True)] == ["111111111111", "222222222222"]

    def test_the_name_is_trimmed(self, ledger):
        excel_source.add_lifecycle("  观察  ", "teal")
        assert LifecycleTag("观察", "teal") in excel_source.load_lifecycle(force=True)

    def test_twelve_characters_is_fine(self, ledger):
        excel_source.add_lifecycle(THIRTEEN[:12], "blue")
        assert excel_source.load_lifecycle(force=True)[-1].name == THIRTEEN[:12]

    @pytest.mark.parametrize(
        "name, color, message",
        [
            ("", "green", "标签名不能为空。"),
            ("   ", "green", "标签名不能为空。"),
            (THIRTEEN, "green", "标签名最多 12 个字。"),
            ("甲,乙", "green", "标签名里不能有逗号、顿号、分号或竖线。"),
            ("甲、乙", "green", "标签名里不能有逗号、顿号、分号或竖线。"),
            ("甲；乙", "green", "标签名里不能有逗号、顿号、分号或竖线。"),
            ("甲|乙", "green", "标签名里不能有逗号、顿号、分号或竖线。"),
            ("观察", "pink", "请从色块里选一个颜色。"),
            ("观察", "", "请从色块里选一个颜色。"),
            ("正常", "green", "已经有「正常」这个标签了。"),
        ],
    )
    def test_refused(self, ledger, name, color, message):
        before = ledger.stat().st_mtime_ns
        with pytest.raises(ExcelSourceError, match=re.escape(message)):
            excel_source.add_lifecycle(name, color)
        assert ledger.stat().st_mtime_ns == before
        assert tag_sheet(ledger) is None

    def test_the_list_has_a_ceiling(self, ledger, monkeypatch):
        monkeypatch.setattr(excel_source, "MAX_LIFECYCLE_TAGS", 4)
        excel_source.add_lifecycle("观察", "amber")
        with pytest.raises(ExcelSourceError, match="标签最多 4 个"):
            excel_source.add_lifecycle("续约", "blue")


class TestRemoveLifecycle:
    def test_takes_it_off_the_list_and_every_account(self, ledger):
        rewrite_ledger(ledger, LIFECYCLE=["正常,结算", "结算"])
        note = excel_source.remove_lifecycle("结算", actor="tester")
        assert note == "删除生命周期标签「结算」（2 个账号上的一起去掉）"
        assert tag_sheet(ledger) == [["正常", "green"], ["风控", "red"]]
        first, second = load_accounts(force=True)
        assert (first.lifecycle, second.lifecycle) == (("正常",), ())
        assert life_cell(ledger, 1) is None
        assert f"tester\t{note}" in audit(ledger)
        assert backups(ledger) == []

    def test_accounts_without_it_are_left_alone(self, ledger):
        rewrite_ledger(ledger, LIFECYCLE=["正常", None])
        assert excel_source.remove_lifecycle("风控") == "删除生命周期标签「风控」"
        assert life_cell(ledger) == "正常"

    def test_the_name_is_trimmed(self, ledger):
        excel_source.remove_lifecycle(" 风控 ")
        assert [tag.name for tag in excel_source.load_lifecycle(force=True)] == ["正常", "结算"]

    def test_an_unknown_name_changes_nothing(self, ledger):
        before = ledger.stat().st_mtime_ns
        assert excel_source.remove_lifecycle("外星人") == ""
        assert ledger.stat().st_mtime_ns == before
        assert tag_sheet(ledger) is None                              # 连表都没建

    def test_an_old_ledger_without_the_column(self, ledger):
        assert excel_source.remove_lifecycle("结算") == "删除生命周期标签「结算」"
        assert "LIFECYCLE" not in raw_rows(ledger)[0]


# ------------------------------------------------------------------ 表格里的小弹层：/accounts/lifecycle
class TestInlineEdit:
    """表格里点铅笔改一个账号的标签：页面脚本 fetch 调 /accounts/lifecycle，回 JSON，表格原地更新。"""

    def test_saves_and_answers_with_the_colors(self, admin, ledger):
        response = set_life(admin, by_account("111111111111").key, "正常", "风控")
        assert response.status_code == 200
        # 结果和别的操作一样从右上角说：粗体一句话、灰字是哪个账号、下面一行现在是什么
        assert response.get_json() == {
            "ok": True, "tone": "ok", "title": "已更新生命周期", "sub": "acct-one@example.com",
            "text": "现在是：正常、风控。",
            "tags": [{"name": "正常", "color": GREEN}, {"name": "风控", "color": RED}],
        }
        assert by_account("111111111111").lifecycle == ("正常", "风控")
        assert life_cell(ledger) == "正常,风控"

    def test_the_same_tags_again(self, admin, ledger):
        key = by_account("111111111111").key
        set_life(admin, key, "正常")
        before = ledger.stat().st_mtime_ns
        assert set_life(admin, key, "正常").get_json() == {
            "ok": True, "tone": "info", "title": "没有改动", "sub": "acct-one@example.com", "text": "",
            "tags": [{"name": "正常", "color": GREEN}],
        }
        assert ledger.stat().st_mtime_ns == before

    def test_clearing_every_tag(self, admin, ledger):
        key = by_account("111111111111").key
        set_life(admin, key, "正常")
        result = set_life(admin, key).get_json()
        assert result["ok"] is True and result["tags"] == []
        assert result["text"] == "现在没有标签。"
        assert by_account("111111111111").lifecycle == () and life_cell(ledger) is None

    def test_an_unknown_tag_is_refused(self, admin, ledger):
        before = ledger.stat().st_mtime_ns
        response = set_life(admin, by_account("111111111111").key, "正常", "外星人")
        assert response.status_code == 400
        assert response.get_json() == {
            "ok": False, "tone": "error", "title": "生命周期没有保存", "sub": "acct-one@example.com",
            "text": "标签清单里没有「外星人」，刷新页面再选。",
        }
        assert ledger.stat().st_mtime_ns == before

    def test_a_tag_already_on_the_account_may_stay(self, admin, ledger):
        """手改 Excel 留下的老标签（清单里没有）：这个账号上本来就有，保存时可以留着，按灰色画。"""
        rewrite_ledger(ledger, **emails(LIFECYCLE=["正常,老标签", None]))
        result = set_life(admin, by_account("111111111111").key, "老标签", "风控").get_json()
        assert result["ok"] is True
        assert result["tags"] == [{"name": "老标签", "color": GRAY}, {"name": "风控", "color": RED}]
        assert by_account("111111111111").lifecycle == ("老标签", "风控")

    def test_but_it_cannot_spread_to_another_account(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["老标签", None]))
        response = set_life(admin, by_account("222222222222").key, "老标签")
        assert response.status_code == 400 and "「老标签」" in response.get_json()["text"]
        assert by_account("222222222222").lifecycle == ()

    def test_a_vanished_account_is_a_conflict(self, admin, ledger):
        response = set_life(admin, "999999999999#9", "正常")
        assert response.status_code == 409
        assert response.get_json() == {
            "ok": False, "tone": "error", "title": "生命周期没有保存", "sub": "",
            "text": "这个账号已经不在台账里了，页面可能已过期，刷新后再试。",
        }

    def test_an_unreadable_ledger_is_a_server_error(self, admin, ledger, monkeypatch):
        csrf = token(admin)                                           # 台账还在的时候先拿到令牌
        monkeypatch.setattr(config, "EXCEL_PATH", ledger.parent / "missing.xlsx")
        response = admin.post("/accounts/lifecycle", data={"csrf": csrf, "key": KEY_ONE, "lifecycle": ["正常"]},
                              headers=FETCH)
        assert response.status_code == 500
        result = response.get_json()
        assert result["ok"] is False and "找不到账号台账文件" in result["text"]

    def test_the_other_columns_are_left_alone(self, admin, ledger):
        before = by_account("111111111111")
        set_life(admin, before.key, "风控")
        after = by_account("111111111111")
        for name in ("ak", "sk", "email", "budget", "enabled", "tg_chat_ids", "start_date"):
            assert getattr(after, name) == getattr(before, name), name

    def test_a_disabled_account_can_be_tagged_too(self, admin, ledger):
        post(admin, "/accounts/toggle", key=by_account("111111111111").key, enabled="0")
        assert set_life(admin, by_account("111111111111").key, "结算").get_json()["ok"] is True
        assert by_account("111111111111").lifecycle == ("结算",)

    def test_without_javascript_it_is_a_normal_post(self, admin, ledger):
        response = set_life(admin, by_account("111111111111").key, "结算", fetch=False)
        assert response.status_code == 302 and response.headers["Location"].endswith("/accounts/")
        assert toasts(page(admin)) == [("ok", "已更新生命周期 acct-one@example.com 现在是：结算。")]

    def test_without_javascript_an_error_is_a_toast(self, admin, ledger):
        response = set_life(admin, by_account("111111111111").key, "外星人", fetch=False)
        assert response.status_code == 302
        assert toasts(page(admin)) == [
            ("error", "生命周期没有保存 acct-one@example.com 标签清单里没有「外星人」，刷新页面再选。"),
        ]

    def test_needs_csrf(self, admin, ledger):
        before = ledger.stat().st_mtime_ns
        response = admin.post("/accounts/lifecycle", data={"key": by_account("111111111111").key, "lifecycle": ["正常"]},
                              headers=FETCH)
        assert response.status_code == 302                           # 和其他 POST 一样：提示后退回
        assert by_account("111111111111").lifecycle == ()
        assert ledger.stat().st_mtime_ns == before

    def test_needs_login(self, client, ledger):
        response = client.post("/accounts/lifecycle", data={"key": KEY_ONE, "lifecycle": ["正常"]}, headers=FETCH)
        assert response.status_code == 302 and "/login" in response.headers["Location"]
        assert life_cell(ledger) is None


# ------------------------------------------------------------------ 「生命周期标签」弹窗：增删清单
class TestTagList:
    """增删完回到这一页，?tags=1 让「生命周期标签」弹窗重新打开，方便接着改。"""

    def test_adding_reopens_the_dialog(self, admin, ledger):
        response = post(admin, "/accounts/lifecycle/add", name="观察", color="amber")
        assert response.status_code == 302 and response.headers["Location"].endswith("/accounts/?tags=1")
        html = admin.get(response.headers["Location"]).get_data(as_text=True)
        assert reopened_dialogs(html) == ["dlg-life"]
        assert toasts(html) == [("ok", "已添加标签「观察」")]
        assert listed_tags(html) == {"正常": 0, "结算": 0, "风控": 0, "观察": 0}

    def test_the_new_tag_is_ready_to_use(self, admin, ledger):
        post(admin, "/accounts/lifecycle/add", name="观察", color="amber")
        result = set_life(admin, by_account("111111111111").key, "观察").get_json()
        assert result["tags"] == [{"name": "观察", "color": AMBER}]
        html = page(admin)
        assert offered_tags(_dialog(html, 'id="dlg-create"')) == ["正常", "结算", "风控", "观察"]
        assert "观察" in filter_chips(html)

    @pytest.mark.parametrize(
        "name, color, message",
        [
            ("正常", "green", "已经有「正常」这个标签了。"),
            ("", "amber", "标签名不能为空。"),
            ("甲,乙", "amber", "标签名里不能有逗号、顿号、分号或竖线。"),
            ("观察", "", "请从色块里选一个颜色。"),
        ],
    )
    def test_a_refused_tag_reopens_the_dialog_with_the_reason(self, admin, ledger, name, color, message):
        before = ledger.stat().st_mtime_ns
        response = post(admin, "/accounts/lifecycle/add", name=name, color=color)
        assert response.status_code == 400
        html = response.get_data(as_text=True)
        assert reopened_dialogs(html) == ["dlg-life"]
        assert life_error(html) == message
        # 填过的原样回填，不用重打
        form = add_form(html)
        assert input_attrs(form, "name")["value"] == name
        assert [field["value"] for field in inputs(form, "color") if field.get("checked")] == ([color] if color else [])
        assert ledger.stat().st_mtime_ns == before

    def test_the_list_has_a_ceiling(self, admin, ledger, monkeypatch):
        monkeypatch.setattr(excel_source, "MAX_LIFECYCLE_TAGS", 3)
        response = post(admin, "/accounts/lifecycle/add", name="观察", color="amber")
        assert response.status_code == 400
        assert life_error(response.get_data(as_text=True)) == "标签最多 3 个。"

    def test_removing_takes_it_off_every_account(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["正常,结算", "结算"]))
        response = post(admin, "/accounts/lifecycle/remove", name="结算")
        assert response.status_code == 302 and response.headers["Location"].endswith("/accounts/?tags=1")
        html = admin.get(response.headers["Location"]).get_data(as_text=True)
        assert reopened_dialogs(html) == ["dlg-life"]
        assert toasts(html) == [("ok", "已删除标签「结算」 2 个账号上的一起去掉了。")]
        assert listed_tags(html) == {"正常": 1, "风控": 0}
        assert (by_account("111111111111").lifecycle, by_account("222222222222").lifecycle) == (("正常",), ())

    def test_removing_an_unknown_tag_is_just_a_note(self, admin, ledger):
        before = ledger.stat().st_mtime_ns
        response = post(admin, "/accounts/lifecycle/remove", name="外星人")
        assert response.status_code == 302 and response.headers["Location"].endswith("/accounts/?tags=1")
        assert toasts(admin.get(response.headers["Location"]).get_data(as_text=True)) == [
            ("info", "没有改动 标签清单里已经没有「外星人」了。"),
        ]
        assert ledger.stat().st_mtime_ns == before

    def test_add_and_remove_need_csrf(self, admin, ledger):
        before = ledger.stat().st_mtime_ns
        for path, data in (("/accounts/lifecycle/add", {"name": "观察", "color": "amber"}),
                           ("/accounts/lifecycle/remove", {"name": "正常"})):
            assert admin.post(path, data=data).status_code == 302, path
        assert excel_source.load_lifecycle(force=True) == DEFAULTS
        assert ledger.stat().st_mtime_ns == before

    def test_add_and_remove_need_login(self, client, ledger):
        for path, data in (("/accounts/lifecycle/add", {"name": "观察", "color": "amber"}),
                           ("/accounts/lifecycle/remove", {"name": "正常"})):
            response = client.post(path, data=data)
            assert response.status_code == 302 and "/login" in response.headers["Location"], path
        assert tag_sheet(ledger) is None


# ------------------------------------------------------------------ 页面
class TestOnThePage:
    def test_the_header_button_opens_the_tag_dialog(self, admin, ledger):
        html = page(admin)
        assert 'data-open="dlg-life"' in html and html.count('id="dlg-life"') == 1
        assert reopened_dialogs(html) == []
        assert reopened_dialogs(admin.get("/accounts/?tags=1").get_data(as_text=True)) == ["dlg-life"]

    def test_the_dialog_lists_each_tag_with_a_count_and_a_two_step_delete(self, admin, ledger):
        _edit(admin, by_account("111111111111"), lifecycle=["正常", "风控"])
        _edit(admin, by_account("222222222222"), lifecycle=["正常"])
        html = page(admin)
        assert listed_tags(html) == {"正常": 2, "结算": 0, "风控": 1}
        rows = re.findall(r'<li class="life-row">(.*?)</li>', tag_dialog(html), re.S)
        for row, name in zip(rows, ["正常", "结算", "风控"]):
            assert 'action="/accounts/lifecycle/remove"' in row and "data-arm" in row   # 点两下才删
            assert input_attrs(row, "name")["value"] == name
            assert input_attrs(row, "csrf")["value"]

    def test_the_add_form_offers_every_color_and_preselects_an_unused_one(self, admin, ledger):
        form = add_form(page(admin))
        colors = inputs(form, "color")
        assert [field["value"] for field in colors] == list(excel_source.LIFECYCLE_COLORS)
        # 绿、灰、红已经被默认的三个占了，新标签默认给下一个没人用的
        assert [field["value"] for field in colors if field.get("checked")] == ["clay"]
        assert input_attrs(form, "name")["maxlength"] == str(excel_source.MAX_LIFECYCLE_NAME)

    def test_one_shared_popover_for_the_whole_table(self, admin, ledger):
        html = page(admin)
        assert html.count('id="life-pop"') == 1
        pop = html[html.index('id="life-pop"') :]
        pop = pop[: pop.index("</form>")]
        assert 'action="/accounts/lifecycle"' in pop and "data-fetch" in pop
        assert input_attrs(pop, "key")["value"] == ""                 # 点哪一行，页面脚本把那一行的 key 填进来
        assert input_attrs(pop, "csrf")["value"]
        assert offered_tags(pop) == ["正常", "结算", "风控"] and checked_tags(pop) == []

    def test_a_stale_tag_shows_in_the_table_and_the_filters_but_not_in_the_list(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["老标签", None]))
        html = page(admin)
        (_, cells), _ = table_rows(html)
        assert text(cells[3]) == "老标签" and f"--tag: {GRAY}" in cells[3]
        assert filter_chips(html) == ["", "正常", "结算", "风控", "老标签", "__none__"]
        assert "老标签" not in listed_tags(html)                     # 清单里没有它，也就没有删除按钮

    def test_a_new_account_starts_as_normal(self, admin, ledger):
        dialog = _dialog(page(admin), 'id="dlg-create"')
        assert offered_tags(dialog) == ["正常", "结算", "风控"]
        assert checked_tags(dialog) == ["正常"]

    def test_unless_the_list_has_no_normal(self, admin, ledger):
        post(admin, "/accounts/lifecycle/remove", name="正常")
        dialog = _dialog(page(admin), 'id="dlg-create"')
        assert offered_tags(dialog) == ["结算", "风控"] and checked_tags(dialog) == []

    def test_the_edit_dialog_ticks_the_accounts_tags(self, admin, ledger):
        _edit(admin, by_account("111111111111"), lifecycle=["结算", "风控"])
        html = page(admin)
        assert checked_tags(_dialog(html, 'id="dlg-edit-1"')) == ["结算", "风控"]
        assert checked_tags(_dialog(html, 'id="dlg-edit-2"')) == []

    def test_a_stale_tag_is_offered_ticked_in_the_edit_dialog(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["正常,老标签", None]))
        dialog = _dialog(page(admin), 'id="dlg-edit-1"')
        assert offered_tags(dialog) == ["正常", "结算", "风控", "老标签"]
        assert checked_tags(dialog) == ["正常", "老标签"]
        assert re.search(r'class="life-check is-stale"[^>]*>\s*<input[^>]*value="老标签"', dialog)

    def test_an_empty_list_offers_nothing(self, admin, ledger):
        for name in ("正常", "结算", "风控"):
            post(admin, "/accounts/lifecycle/remove", name=name)
        html = page(admin)
        assert excel_source.load_lifecycle(force=True) == []
        assert offered_tags(_dialog(html, 'id="dlg-create"')) == []
        assert filter_chips(html) == ["", "__none__"]


# ------------------------------------------------------------------ 新增 / 修改弹窗里的多选框
class TestInTheDialogs:
    """新增、修改都只认清单里的标签；修改时再加上这个账号身上本来就有的（手改 Excel 的老标签）。"""

    def test_a_new_account_gets_the_ticked_tags(self, admin, ledger):
        post(admin, "/accounts/create", **NEW_FORM, lifecycle=["结算", "风控"])
        assert by_account("333333333333").lifecycle == ("结算", "风控")
        assert life_cell(ledger, 2) == "结算,风控"

    def test_a_new_account_without_ticks(self, admin, ledger):
        assert post(admin, "/accounts/create", **NEW_FORM).status_code == 302
        assert by_account("333333333333").lifecycle == ()

    def test_an_unknown_tag_is_refused_on_create(self, admin, ledger):
        before = ledger.stat().st_mtime_ns
        response = post(admin, "/accounts/create", **NEW_FORM, lifecycle=["正常", "外星人"])
        assert response.status_code == 400
        assert "生命周期里没有「外星人」，先在标签清单里加上。" in response.get_data(as_text=True)
        assert ledger.stat().st_mtime_ns == before

    def test_another_accounts_stale_tag_is_unknown_on_create(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["老标签", None]))
        response = post(admin, "/accounts/create", **NEW_FORM, lifecycle=["老标签"])
        assert response.status_code == 400
        assert by_account("333333333333") is None

    def test_an_unknown_tag_is_refused_on_edit(self, admin, ledger):
        response = _edit(admin, by_account("111111111111"), lifecycle=["外星人"], budget="777")
        assert response.status_code == 400
        assert "生命周期里没有「外星人」" in response.get_data(as_text=True)
        assert by_account("111111111111").budget == 500000

    def test_a_stale_tag_already_on_the_account_is_kept(self, admin, ledger):
        """手改过的老标签别一保存就丢，也别让整张表单存不进去。"""
        rewrite_ledger(ledger, **emails(LIFECYCLE=["正常,老标签", None]))
        response = _edit(admin, by_account("111111111111"), budget="777")     # 原样勾着
        assert response.status_code == 302
        after = by_account("111111111111")
        assert after.budget == 777 and after.lifecycle == ("正常", "老标签")

    def test_a_stale_tag_can_be_dropped(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["正常,老标签", None]))
        _edit(admin, by_account("111111111111"), lifecycle=["正常"])
        assert by_account("111111111111").lifecycle == ("正常",)

    def test_a_stale_tag_cannot_spread_through_another_accounts_dialog(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["老标签", None]))
        response = _edit(admin, by_account("222222222222"), lifecycle=["老标签"])
        assert response.status_code == 400
        assert by_account("222222222222").lifecycle == ()

    def test_a_failed_create_keeps_the_ticks(self, admin, ledger):
        html = post(admin, "/accounts/create", **{**NEW_FORM, "budget": "abc"},
                    lifecycle=["结算", "风控"]).get_data(as_text=True)
        assert checked_tags(_dialog(html, 'id="dlg-create"')) == ["结算", "风控"]

    def test_a_failed_create_with_nothing_ticked_stays_empty(self, admin, ledger):
        """「正常」只是新弹窗的默认勾选；用户自己取消了，回填时不能又勾回去。"""
        html = post(admin, "/accounts/create", **{**NEW_FORM, "budget": "abc"}).get_data(as_text=True)
        assert checked_tags(_dialog(html, 'id="dlg-create"')) == []

    def test_a_failed_edit_keeps_the_ticks_including_a_stale_one(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["老标签", None]))
        html = _edit(admin, by_account("111111111111"), lifecycle=["老标签", "风控"],
                     budget="abc").get_data(as_text=True)
        dialog = _dialog(html, 'id="dlg-edit-1"')
        assert checked_tags(dialog) == ["风控", "老标签"]             # 清单里的在前，老标签跟在后面
        assert reopened_dialogs(html) == ["dlg-edit-1"]

    def test_the_audit_log_names_the_change(self, admin, ledger):
        _edit(admin, by_account("111111111111"), lifecycle=["结算", "风控"])
        assert "LIFECYCLE 空 → 结算,风控" in audit(ledger)

    def test_the_same_tags_written_differently_are_not_a_change(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=["正常，结算", None]))       # 手打的全角逗号
        before = ledger.stat().st_mtime_ns
        _edit(admin, by_account("111111111111"))
        assert ledger.stat().st_mtime_ns == before

    def test_names_are_at_most_twelve_characters(self):
        _, errors = excel_source.validate({**NEW_FORM, "lifecycle": THIRTEEN}, [], creating=True)
        assert errors == [f"生命周期标签最多 12 个字：「{THIRTEEN}」太长了。"]

    def test_a_long_stale_tag_does_not_block_saving(self, admin, ledger):
        rewrite_ledger(ledger, **emails(LIFECYCLE=[f"正常,{THIRTEEN}", None]))
        response = _edit(admin, by_account("111111111111"), budget="777")
        assert response.status_code == 302
        after = by_account("111111111111")
        assert after.budget == 777 and after.lifecycle == ("正常", THIRTEEN)
